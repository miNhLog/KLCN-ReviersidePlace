from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"C:\Users\KIEULOAN\Documents\Year 4\KLTN\RevirsidePlace")
OUT = Path(r"C:\Users\KIEULOAN\Documents\Intern\Daily-Report\LeMinhLong_BaoCaoNgay_12-10.docx")
IMAGES = [
    Path(r"C:\Users\KIEULOAN\AppData\Local\Temp\codex-clipboard-85fb0891-3f12-4d01-a92c-d57838eca5e4.png"),
    Path(r"C:\Users\KIEULOAN\AppData\Local\Temp\codex-clipboard-a0b886d2-5748-4c91-8fad-8e6cd0a784a5.png"),
    Path(r"C:\Users\KIEULOAN\AppData\Local\Temp\codex-clipboard-df9f20e3-8e40-405f-9b39-383e40201a7e.png"),
    Path(r"C:\Users\KIEULOAN\AppData\Local\Temp\codex-clipboard-4bf0c0cb-3dff-4d02-bd27-dc850599aea0.png"),
]

CAPTIONS = [
    "Hình 1. Giao diện quản lý thực đơn",
    "Hình 2. Chi tiết thực đơn và danh sách món ăn",
    "Hình 3. Giao diện quản lý món ăn",
    "Hình 4. Chi tiết món ăn và các thực đơn đang sử dụng",
]

for image in IMAGES:
    if not image.exists():
        raise FileNotFoundError(image)

doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(2.0)
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.5)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(13)
normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
normal.paragraph_format.space_after = Pt(2)
normal.paragraph_format.line_spacing = 1.08

def set_run_font(run, size=13, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic

def add_label(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(5)
    set_run_font(p.add_run(text), 13, bold=True)
    return p

def add_bullet(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.first_line_indent = Cm(-0.45)
    p.paragraph_format.space_after = Pt(2)
    set_run_font(p.add_run("-  "), 13)
    set_run_font(p.add_run(text), 13)
    return p

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(10)
set_run_font(title.add_run("BÁO CÁO NGÀY 12/10/2026"), 16, bold=True)

project = doc.add_paragraph()
project.paragraph_format.space_after = Pt(6)
set_run_font(project.add_run("Dự án: "), 13, bold=True)
set_run_font(project.add_run("Xây dựng hệ thống quản lý và khuyến nghị đặt tiệc cưới tại Chuỗi trung tâm hội nghị tiệc cưới Riverside Palace."), 13)

add_label("Hôm trước:")
previous = [
    "Hoàn thiện chức năng Nhật ký hệ thống dành cho Quản trị viên, hỗ trợ theo dõi các thao tác quản trị tài khoản, phân quyền và bảo mật.",
    "Bổ sung bộ lọc nhật ký theo loại thao tác, thời gian, nội dung tìm kiếm và phân trang; giới hạn dữ liệu theo đúng phạm vi quản trị hệ thống.",
    "Xây dựng quy trình gửi, theo dõi, duyệt và từ chối yêu cầu thay đổi vai trò nhân viên giữa Nhân viên điều phối và Quản lý sảnh.",
    "Bổ sung kiểm tra phân công hiện tại trước khi duyệt thay đổi vai trò nhằm bảo đảm tính nhất quán của dữ liệu nhân sự.",
]
for item in previous:
    add_bullet(item)

add_label("Hôm nay:")
today = [
    "Tách màn hình Thực đơn và Món ăn thành hai chức năng quản lý độc lập dành cho vai trò Quản lý.",
    "Hoàn thiện chức năng Quản lý thực đơn, gồm thống kê tổng quan, tìm kiếm, lọc trạng thái, phân trang phía máy chủ, thêm mới, chỉnh sửa và cập nhật trạng thái thực đơn.",
    "Bổ sung giao diện xem chi tiết thực đơn ngay trong danh sách; hỗ trợ thêm nhiều món, bỏ món và sắp xếp thứ tự món bằng thao tác kéo thả.",
    "Bổ sung cơ chế khóa thay đổi thành phần món đối với thực đơn đã được sử dụng trong đặt tiệc, đồng thời vẫn cho phép cập nhật tên, mô tả, giá và trạng thái.",
    "Hoàn thiện chức năng Quản lý món ăn theo dạng lưới thẻ, gồm thống kê, tìm kiếm, lọc danh mục, lọc trạng thái và phân trang phía máy chủ.",
    "Bổ sung chức năng thêm, chỉnh sửa món ăn, cập nhật giá, tải lên, thay thế hoặc xóa hình ảnh và chuyển đổi trạng thái Đang phục vụ hoặc Tạm ngưng.",
    "Xây dựng panel chi tiết món ăn, hiển thị số lượng và danh sách các thực đơn đang sử dụng món, bao gồm cả thực đơn tiêu chuẩn và thực đơn tùy chỉnh.",
    "Hoàn thiện kiểm tra an toàn tệp ảnh theo dung lượng, phần mở rộng, MIME type và chữ ký tệp; giữ nguyên quan hệ món trong thực đơn khi món được tạm ngưng.",
]
for item in today:
    add_bullet(item)

add_label("Kết quả:")
doc.add_page_break()

for index, (image_path, caption) in enumerate(zip(IMAGES, CAPTIONS)):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(str(image_path), width=Cm(16.0))

    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(10)
    set_run_font(cp.add_run(caption), 13, italic=True)

    if index == 1:
        doc.add_page_break()

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
