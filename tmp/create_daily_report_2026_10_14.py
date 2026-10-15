from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path(r"C:\Users\KIEULOAN\Documents\Intern\Daily-Report\LeMinhLong_BaoCaoNgay_14-10-2026.docx")


def set_run_font(run, size=14, bold=False):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.bold = bold


def add_label(doc, text):
    p = doc.add_paragraph(style="Report Label")
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(1)
    set_run_font(p.add_run(text), size=14, bold=True)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="Report Bullet")
    p.paragraph_format.left_indent = Cm(1.05)
    p.paragraph_format.first_line_indent = Cm(-0.55)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    set_run_font(p.add_run("-\t" + text), size=14)
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(2.0)
section.left_margin = Cm(2.4)
section.right_margin = Cm(2.2)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
normal.font.size = Pt(14)

title_style_ppr = doc.styles["Title"]._element.get_or_add_pPr()
title_style_border = title_style_ppr.find(qn("w:pBdr"))
if title_style_border is not None:
    title_style_ppr.remove(title_style_border)

for name in ("Report Label", "Report Bullet"):
    if name not in doc.styles:
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = normal

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(7)
title.paragraph_format.line_spacing = 1.0
title_ppr = title._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)
set_run_font(title.add_run("BÁO CÁO NGÀY 14/10/2026"), size=16, bold=True)

project = doc.add_paragraph()
project.paragraph_format.space_after = Pt(2)
project.paragraph_format.line_spacing = 1.0
set_run_font(project.add_run("Dự án: "), size=14, bold=True)
set_run_font(project.add_run("Xây dựng hệ thống quản lý và khuyến nghị đặt tiệc cưới."), size=14, bold=True)

add_label(doc, "Hôm trước:")
for item in [
    "Hoàn thiện chức năng Hợp đồng & Thanh toán dành cho vai trò Quản lý, gồm hai tab Hợp đồng và Thanh toán trên cùng một màn hình quản lý.",
    "Xây dựng quy trình tạo hợp đồng từ Booking đã xác nhận; hệ thống tự sinh mã hợp đồng, tạo hợp đồng ở trạng thái Bản nháp và điền sẵn tổng giá trị, nội dung, điều khoản thanh toán.",
    "Bổ sung chức năng chỉnh sửa hợp đồng ở trạng thái Bản nháp, kích hoạt hợp đồng sau khi kiểm tra đầy đủ thông tin, hoàn tất khi đã thanh toán đủ và hủy hợp đồng theo đúng luồng trạng thái.",
    "Bổ sung danh sách hợp đồng với thống kê tổng quan, tìm kiếm, lọc trạng thái, lọc theo thời gian, phân trang phía máy chủ và giao diện xem chi tiết hợp đồng bằng khung thông tin bên phải.",
]:
    add_bullet(doc, item)

add_label(doc, "Hôm nay:")
for item in [
    "Hoàn thiện luồng Hợp đồng & Thanh toán ở cả API và giao diện Quản lý; bổ sung tạo, chỉnh sửa, kích hoạt, hoàn tất, hủy hợp đồng và ghi nhận thanh toán theo đúng quy tắc nghiệp vụ.",
    "Bổ sung kiểm tra điều kiện hợp đồng, chính sách thanh toán, tự sinh mã hợp đồng, nhật ký thao tác và các bài kiểm thử cho quy tắc nghiệp vụ hợp đồng.",
    "Xây dựng chức năng tạo file PDF hợp đồng từ mẫu, đồng thời cập nhật màn hình đặt tiệc để mở nhanh quy trình tạo và theo dõi hợp đồng.",
    "Chuẩn hóa tiêu đề, breadcrumb và nút tạo mới trên các màn hình quản trị để giao diện giữa các chức năng thống nhất và hiển thị tốt trên thiết bị di động.",
    "Bổ sung bộ sưu tập ảnh cho trang chi tiết sảnh, gồm ảnh thu nhỏ, chuyển ảnh trước hoặc sau, xem toàn màn hình và dữ liệu nhiều ảnh cho từng sảnh.",
]:
    add_bullet(doc, item)

add_label(doc, "Kết quả:")
add_bullet(doc, "Hoàn thành chức năng Hợp đồng & Thanh toán cho vai trò Quản lý; đồng thời cải thiện tính nhất quán của giao diện quản trị và trải nghiệm xem hình ảnh sảnh.")

props = doc.core_properties
props.title = "Báo cáo ngày 14 10 2026"
props.subject = "Báo cáo công việc hằng ngày dự án Riverside Palace"
props.author = "Lê Minh Long"

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
