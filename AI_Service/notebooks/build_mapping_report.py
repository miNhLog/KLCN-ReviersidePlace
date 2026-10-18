"""Sinh bao cao doi chieu de cuong <-> ma nguon tu metrics_summary.json.

Xuat 3 dinh dang trong notebooks/: recommender_mapping.md / .html / .docx
Chay: python -X utf8 notebooks/build_mapping_report.py  (workdir: AI_Service)
"""
import base64
import json
from pathlib import Path

HERE = Path(__file__).resolve().parent
PLOT_DIR = HERE / "output_plots"
METRICS = PLOT_DIR / "metrics_summary.json"
TITLE = "Hệ thống gợi ý Riverside Palace — Đối chiếu đề cương CNTT-KLCN257 & mã nguồn"

MODELS = ["WSM", "CBF", "CF", "Hybrid"]


def f4(v) -> str:
    try:
        return f"{float(v):.4f}"
    except (TypeError, ValueError):
        return str(v)


def load_metrics() -> dict:
    if not METRICS.exists():
        raise SystemExit(f"Chua co {METRICS} — chay truoc: python -X utf8 notebooks/export_metrics.py")
    return json.loads(METRICS.read_text(encoding="utf-8"))


def build_sections(d: dict) -> list:
    rec = d["recommender_comparison"]
    wsm = d["wsm_tuning"]
    hw = d["hybrid_weights"]
    meta = d["meta"]
    w = wsm["weights"]
    sections = []

    sections.append({
        "heading": "1. Thuật toán & vị trí mã nguồn",
        "paras": ["Toàn bộ thuật toán khuyến nghị nằm trong AI_Service/app/core (Python). "
                  "Mô hình chính là Hybrid (kết hợp CF + CBF + MCDM/WSM)."],
        "table": {
            "headers": ["Thuật toán", "File", "Hàm (dòng)", "Vai trò"],
            "rows": [
                ["CBF — Content-Based Filtering",
                 "AI_Service/app/core/cbf_engine.py",
                 "build_index() :73 · cbf_score() :95",
                 "Vector hóa TF-IDF hồ sơ hall/menu/decor, so khớp cosine với sở thích"],
                ["CF — Collaborative Filtering",
                 "AI_Service/app/core/cf_engine.py",
                 "fit() :79 · predict_rating() :108 · cf_score() :123",
                 "Ma trận user–item (SampleId × Hall_Menu_Decor), phân rã TruncatedSVD n=20"],
                ["MCDM / WSM",
                 "AI_Service/app/core/wsm_engine.py",
                 "score_candidate() :72 · rank_top3() :142",
                 "Tổng có trọng số w1..w4: ngân sách, sức chứa, phong cách, dịch vụ"],
                ["Hybrid — mô hình chính",
                 "AI_Service/app/core/hybrid_engine.py",
                 "load_hybrid_weights() :24 · fuse() :55 · rank_candidates() :62",
                 "s_final = α·s_cf + β·s_cbf + γ·S_wsm (α, β, γ học từ dữ liệu)"],
                ["LLM + RAG — trợ lý AI",
                 "AI_Service/app/core/rag_engine.py",
                 "chat() ~:586 · extract_entities() :794",
                 "Trợ lý tiếng Việt (Gemini gemini-3.5-flash-lite) + trích thực thể cho khảo sát"],
            ],
        },
    })

    sections.append({
        "heading": "2. Đối chiếu yêu cầu đề cương",
        "paras": [],
        "table": {
            "headers": ["Yêu cầu đề cương", "Hiện trạng", "Trạng thái"],
            "rows": [
                ["Môi trường LLM và RAG",
                 "rag_engine.py (retrieval + Gemini gemini-3.5-flash-lite); kho tri thức riverside_knowledge.md",
                 "Đạt ✅"],
                ["Cài đặt bằng Python + FastAPI, tích hợp ASP.NET Core",
                 "AI_Service (FastAPI :8000); Backend/Frontend gọi qua API nội bộ",
                 "Đạt ✅"],
                ["Content-Based Filtering (CBF)",
                 "cbf_engine.py — TF-IDF + cosine",
                 "Đạt ✅"],
                ["Collaborative Filtering (CF) trên dữ liệu tiệc cưới trước đó",
                 "cf_engine.py (SVD) đã hoạt động; hiện huấn luyện trên dữ liệu SYNTHETIC (wedding_dataset_5k.csv)",
                 "Một phần ⚠️ — cần dữ liệu thật"],
                ["Hybrid",
                 "hybrid_engine.py; α/β/γ học từ dữ liệu (hybrid_weights.json)",
                 "Đạt ✅"],
                ["MCDM / WSM",
                 "wsm_engine.py; w1..w4 tối ưu bằng SLSQP (trained_weights.json)",
                 "Đạt ✅"],
                ["Phân tích đầu vào: ngân sách, số bàn, thời gian, phong cách",
                 "wsm_engine.score_candidate + rag_engine.extract_entities (ngày/ca/khách)",
                 "Đạt ✅"],
                ["Gợi ý Top-3 sảnh / menu / decor",
                 "rank_top3() và rank_candidates(); hiển thị tại /Recommendation/Index",
                 "Đạt ✅"],
                ["Docker & Docker Compose",
                 "docker-compose.yml (service ai) + Dockerfile từng tầng",
                 "Đạt ✅"],
                ["Jupyter Notebook thực nghiệm",
                 "notebooks/demo_recommender.ipynb + các script thực nghiệm + biểu đồ output_plots",
                 "Đạt ✅"],
                ["Tài liệu tham khảo (Aggarwal, Koren, Lewis, Ricci, scikit-learn)",
                 "Áp dụng trong cbf_engine / cf_engine / rag_engine / notebook demo",
                 "Đạt ✅"],
                ["Dữ liệu các tiệc cưới thành công thật",
                 "wedding_reviews_real.csv hiện chỉ là mẫu mô phỏng (SAMPLE_NOTE); cần export từ DB thật",
                 "Một phần ⚠️"],
            ],
        },
    })

    abl = hw.get("ablation") or {}
    abl_txt = ""
    if abl:
        label = {"alpha=0": "α=0 (bỏ CF)", "beta=0": "β=0 (bỏ CBF)", "gamma=0": "γ=0 (bỏ WSM)"}
        abl_txt = ("Ablation — bỏ từng thành phần (MSE trên tập validation): "
                   + "; ".join(f"{label.get(k, k)}: {f4(v.get('mse_val'))}" for k, v in abl.items()) + ".")
    para_hybrid = (
        f"Mô hình chính: Hybrid — s_final = α·s_cf + β·s_cbf + γ·S_wsm với "
        f"α={float(hw['alpha']):.2f}, β={float(hw['beta']):.2f}, γ={float(hw['gamma']):.2f} "
        f"(học từ dữ liệu bằng SLSQP + grid-search, lưu tại app/core/hybrid_weights.json; "
        f"MSE train={f4(hw.get('mse_train'))}, MSE val={f4(hw.get('mse_val'))})."
    )
    dep = d.get("deployed_weights") or {}
    para_deployed = ""
    if dep:
        para_deployed = (
            f"Trọng số WSM đang triển khai (wsm_engine dùng để xếp hạng thực tế): "
            f"w1={f4(dep.get('w1_budget'))}, w2={f4(dep.get('w2_capacity'))}, "
            f"w3={f4(dep.get('w3_style'))}, w4={f4(dep.get('w4_service'))}. "
            "Bảng so sánh 4 mô hình bên dưới đánh giá theo bộ trọng số này."
        )
    para_wsm = (
        f"Thực nghiệm huấn luyện lại trọng số WSM bằng SLSQP (min MSE giữa điểm S và Satisfaction/5.0; "
        f"kết quả chỉ để đối chiếu, KHÔNG ghi đè trọng số đang triển khai) trên "
        f"{wsm['n_train']} mẫu huấn luyện / {wsm['n_test']} mẫu kiểm thử. Baseline w=[0.25, 0.25, 0.25, 0.25] "
        f"so với tối ưu w=[{f4(w['w1_budget'])}, {f4(w['w2_capacity'])}, {f4(w['w3_style'])}, {f4(w['w4_service'])}]. "
        f"MSE test: baseline={f4(wsm['mse_test']['baseline'])} → tối ưu={f4(wsm['mse_test']['optimized'])}."
    )
    sections.append({
        "heading": "3. Mô hình chính & so sánh mô hình",
        "paras": [
            para_hybrid,
            abl_txt if abl_txt else "",
            para_deployed,
            para_wsm,
            (f"So sánh 4 mô hình (dùng trọng số WSM đang triển khai) trên cùng tập test "
             f"(Top-3, relevant = grade ≥ 2, {meta['n_eval']} mẫu × {meta['n_candidates']} ứng viên, "
             f"seed {meta['seed']}):"),
        ],
        "table": {
            "headers": ["Mô hình", "P@3", "R@3", "NDCG@3", "HitRate@3"],
            "rows": [[m, f4(rec[m]["p3"]), f4(rec[m]["r3"]), f4(rec[m]["ndcg3"]), f4(rec[m]["hit3"])]
                     for m in MODELS],
        },
    })
    sections[-1]["paras"] = [p for p in sections[-1]["paras"] if p]
    sections[-1]["table2"] = {
        "headers": ["Cấu hình WSM", "P@3", "R@3", "NDCG@3", "HitRate@3", "MSE train", "MSE test"],
        "rows": [
            ["Baseline (w = 0.25 đều)",
             f4(wsm["baseline"]["p3"]), f4(wsm["baseline"]["r3"]),
             f4(wsm["baseline"]["ndcg3"]), f4(wsm["baseline"]["hit3"]),
             f4(wsm["mse_train"]["baseline"]), f4(wsm["mse_test"]["baseline"])],
            ["Optimized (SLSQP)",
             f4(wsm["optimized"]["p3"]), f4(wsm["optimized"]["r3"]),
             f4(wsm["optimized"]["ndcg3"]), f4(wsm["optimized"]["hit3"]),
             f4(wsm["mse_train"]["optimized"]), f4(wsm["mse_test"]["optimized"])],
        ],
        "caption": "Baseline vs Optimized (WSM) — cùng tập test Top-3:",
    }

    imgs = [
        ("hybrid_comparison.png", "notebooks/evaluate_hybrid.py:187",
         "So sánh WSM/CBF/CF/Hybrid theo P@3, R@3, NDCG@3, HitRate@3"),
        ("metrics_comparison.png", "notebooks/train_and_evaluate.py:220",
         "Baseline vs Optimized (4 chỉ số Top-3)"),
        ("weight_distribution.png", "notebooks/train_and_evaluate.py:228",
         "Phân bố trọng số WSM tối ưu (w1..w4)"),
        ("score_distribution.png", "notebooks/train_and_evaluate.py:237",
         "Phân phối điểm WSM tối ưu trên tập test"),
    ]
    sections.append({
        "heading": "4. Hình ảnh minh chứng",
        "paras": ["Các biểu đồ được sinh tự động khi chạy thực nghiệm, lưu tại AI_Service/notebooks/output_plots."],
        "table": {
            "headers": ["Ảnh", "Script tạo", "Nội dung"],
            "rows": [[a, s, c] for a, s, c in imgs],
        },
        "images": [{"file": a, "rel": f"output_plots/{a}", "caption": f"Hình: {c} ({a})"} for a, _s, c in imgs],
    })

    sections.append({
        "heading": "5. Ghi chú & hạn chế",
        "paras": [
            f"Nguồn dữ liệu: {d['data_source']}. Kết quả mang tính mô phỏng; khi có dữ liệu tiệc cưới thật, "
            "chạy app/core/export_reviews.py rồi huấn luyện lại CF/Hybrid/WSM.",
            "Toàn bộ chỉ số số học được lưu tự động tại notebooks/output_plots/metrics_summary.json "
            "(sinh bởi notebooks/export_metrics.py).",
            (f"Hệ thống giữ bộ trọng số WSM đang triển khai: w1={f4(dep.get('w1_budget'))}, "
             f"w2={f4(dep.get('w2_capacity'))}, w3={f4(dep.get('w3_style'))}, w4={f4(dep.get('w4_service'))} "
             "để tiêu chí phong cách (w3) vẫn đóng góp vào điểm. Khi huấn luyện lại trên dữ liệu mô phỏng, "
             f"SLSQP đẩy w3→{f4(w['w3_style'])} (không triển khai) — cần huấn luyện lại khi có dữ liệu thật."),
            "Hạn chế: script huấn luyện dùng điểm r1 đơn giản, chưa đồng bộ với chia bậc vượt ngân sách "
            "trong wsm_engine.score_candidate — cần đồng bộ khi huấn luyện với dữ liệu thật.",
        ],
        "table": None,
    })
    return sections


def render_md(sections: list, out: Path):
    lines = [f"# {TITLE}", ""]
    for s in sections:
        lines += [f"## {s['heading']}", ""]
        for p in s.get("paras", []):
            lines += [p, ""]
        for key in ("table", "table2"):
            t = s.get(key)
            if not t:
                continue
            if t.get("caption"):
                lines += [t["caption"], ""]
            lines.append("| " + " | ".join(t["headers"]) + " |")
            lines.append("|" + "|".join(["---"] * len(t["headers"])) + "|")
            for row in t["rows"]:
                lines.append("| " + " | ".join(str(c) for c in row) + " |")
            lines.append("")
        for img in s.get("images", []):
            lines += [f"![{img['caption']}]({img.get('rel', img['file'])})", ""]
    out.write_text("\n".join(lines), encoding="utf-8")


def esc(t) -> str:
    return str(t).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def render_html(sections: list, out: Path):
    css = ("body{font-family:Calibri,Arial,sans-serif;margin:2cm;color:#222}"
           "h1{font-size:20pt}h2{font-size:14pt;margin-top:24px;border-bottom:1px solid #ccc;padding-bottom:4px}"
           "table{border-collapse:collapse;margin:8px 0;font-size:10.5pt}"
           "th,td{border:1px solid #999;padding:5px 8px;text-align:left;vertical-align:top}"
           "th{background:#f0f0f0}img{max-width:100%;border:1px solid #ddd;margin-top:6px}"
           "p{font-size:11pt;line-height:1.5}")
    parts = [f'<!doctype html><html lang="vi"><head><meta charset="utf-8">'
             f"<title>{esc(TITLE)}</title><style>{css}</style></head><body>",
             f"<h1>{esc(TITLE)}</h1>"]
    for s in sections:
        parts.append(f"<h2>{esc(s['heading'])}</h2>")
        for p in s.get("paras", []):
            parts.append(f"<p>{esc(p)}</p>")
        for key in ("table", "table2"):
            t = s.get(key)
            if not t:
                continue
            if t.get("caption"):
                parts.append(f"<p><b>{esc(t['caption'])}</b></p>")
            parts.append("<table><thead><tr>" + "".join(f"<th>{esc(h)}</th>" for h in t["headers"]) + "</tr></thead><tbody>")
            for row in t["rows"]:
                parts.append("<tr>" + "".join(f"<td>{esc(c)}</td>" for c in row) + "</tr>")
            parts.append("</tbody></table>")
        for img in s.get("images", []):
            p = PLOT_DIR / img["file"]
            if p.exists():
                b64 = base64.b64encode(p.read_bytes()).decode("ascii")
                parts.append(f'<img alt="{esc(img["caption"])}" src="data:image/png;base64,{b64}">')
                parts.append(f'<p style="font-size:10pt;color:#555">{esc(img["caption"])}</p>')
    parts.append("</body></html>")
    out.write_text("".join(parts), encoding="utf-8")


def render_docx(sections: list, out: Path):
    from docx import Document
    from docx.shared import Cm

    doc = Document()
    doc.add_heading(TITLE, level=0)
    for s in sections:
        doc.add_heading(s["heading"], level=1)
        for p in s.get("paras", []):
            doc.add_paragraph(p)
        for key in ("table", "table2"):
            t = s.get(key)
            if not t:
                continue
            if t.get("caption"):
                doc.add_paragraph(t["caption"])
            table = doc.add_table(rows=1, cols=len(t["headers"]))
            table.style = "Table Grid"
            for i, h in enumerate(t["headers"]):
                cell = table.rows[0].cells[i]
                cell.text = str(h)
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            for row in t["rows"]:
                cells = table.add_row().cells
                for i, c in enumerate(row):
                    cells[i].text = str(c)
        for img in s.get("images", []):
            p = PLOT_DIR / img["file"]
            if p.exists() and p.stat().st_size > 0:
                doc.add_picture(str(p), width=Cm(15))
                cap = doc.add_paragraph(img["caption"])
                if cap.runs:
                    cap.runs[0].italic = True
    doc.save(str(out))


def main():
    d = load_metrics()
    sections = build_sections(d)
    render_md(sections, HERE / "recommender_mapping.md")
    render_html(sections, HERE / "recommender_mapping.html")
    render_docx(sections, HERE / "recommender_mapping.docx")
    for name in ("recommender_mapping.md", "recommender_mapping.html", "recommender_mapping.docx"):
        p = HERE / name
        print(f"Da luu -> {p} ({p.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
